"""Build the editable test bank and Markdown copies from bank.json.

Run with the Codex bundled Python runtime, which supplies python-docx.
"""
from collections import Counter
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'bank.json').read_text(encoding='utf-8'))
QUESTIONS = DATA['questions']
LETTERS = 'ABCD'


def check():
    ids = [q['id'] for q in QUESTIONS]
    assert len(ids) == len(set(ids)) == 46, 'Expected 46 unique question IDs'
    expected = [f'{prefix}{i:02}' for prefix, n in [('TF', 20), ('MC', 20), ('OE', 6)] for i in range(1, n + 1)]
    assert ids == expected, 'Missing, extra or out-of-order IDs'
    assert Counter(q['topic'] for q in QUESTIONS) == {'Hardware': 16, 'Networks': 16, 'AI Behaviour': 14}
    for q in QUESTIONS:
        assert q['prompt'].strip() and q['source'].strip()
        if q['id'].startswith('MC'):
            assert len(q['options']) == len(set(q['options'])) == 4
            assert q['answer'] in LETTERS and q['explanation'].strip()
        elif q['id'].startswith('TF'):
            assert q['answer'] in ('True', 'False') and q['explanation'].strip()
        else:
            assert len(q['tasks']) == 3 and len(q['key_points']) == 4
            assert len(q['sample'].split()) > 100
    assert Counter(q['answer'] for q in QUESTIONS[:20]) == {'True': 10, 'False': 10}
    assert Counter(q['answer'] for q in QUESTIONS[20:40]) == dict.fromkeys(LETTERS, 5)
    print('Passed: 46 unique IDs, topic counts, 20 TF, 20 MC, 6 guided OE, complete keys and developed sample answers.')


def markdown():
    lines = ['# Hardware Networks and AI Behaviour Test Bank', '', '**Team Doom | Version 1.0 draft | 30 October 2026**', '',
             'Use these questions to practise the main ideas in Hardware, Networks and AI Behaviour. No equipment or external website is needed. Write answers separately, then check AnswerKey.md.', '',
             'The bank is a draft for teammate review. Source-topic labels show where each idea is taught; final Study Guide section references will be added when that guide is available.', '']
    answers = ['# Test Bank Answer Key', '', '**Version 1.0 draft.** Try the questions before reading these answers. For open-ended questions, accept equivalent correct explanations, diagrams or worked steps.', '']
    for prefix, title, instruction in [
        ('TF', 'Part A True or False', 'Write True or False for each statement.'),
        ('MC', 'Part B Multiple Choice', 'Choose ONE answer, A, B, C or D, for each question.'),
        ('OE', 'Part C Explain Your Answer', 'Answer all three parts of each question. Aim for a developed response, normally more than 100 words in total per question, or equivalent reasoning using a labeled diagram, table or step-by-step trace. Keep the explanation simple and use the supplied situation.')]:
        lines += ['## ' + title, '', instruction, '']
        answers += ['## ' + title, '']
        topic = None
        for q in (q for q in QUESTIONS if q['id'].startswith(prefix)):
            if topic != q['topic']:
                topic = q['topic']
                lines += ['### ' + topic, '']
                answers += ['### ' + topic, '']
            lines += [f"#### {q['id']}", '', q['prompt'], '']
            if prefix == 'MC':
                lines += [f'{letter}. {option}<br>' for letter, option in zip(LETTERS, q['options'])] + ['']
            if prefix == 'OE':
                lines += [f'{letter}) {task}' for letter, task in zip('abc', q['tasks'])] + ['']
                answers += [f"#### {q['id']} Sample answer", '', q['sample'], '', '**Expected key points**', ''] + ['- ' + point for point in q['key_points']] + ['']
            else:
                answer = q['answer']
                if prefix == 'MC':
                    answer += ' - ' + q['options'][LETTERS.index(q['answer'])]
                answers += [f"**{q['id']} {answer}.** {q['explanation']}", '']
            lines += [f"*Topic source: {q['source']}.*", '']
    lines += ['## Source Materials', '',
              '- [Hardware](../Learning%20Materials/Hardware/hardware.md)',
              '- [Network Fundamentals](../Learning%20Materials/Network/network_fundamentals.md)',
              '- [Network Beginner Prerequisites](../Learning%20Materials/Network/beginner-prerequisites.md)',
              '- [AI Behaviour](../Learning%20Materials/AIBehaviour/AIBehaviour.md)', '']
    (ROOT / 'Questions.md').write_text('\n'.join(lines), encoding='utf-8')
    (ROOT / 'AnswerKey.md').write_text('\n'.join(answers), encoding='utf-8')


def word():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    # Some bundled default styles carry a blue paragraph rule. Remove it from
    # the export so titles and other text are separated only by whitespace.
    for borders in list(doc.styles.element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.68)
    section.left_margin = section.right_margin = Inches(0.8)
    section.footer_distance = Inches(0.28)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Title'].paragraph_format.space_after = Pt(9)
    doc.styles['Subtitle'].font.size = Pt(11)
    for name, size in [('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Team Doom  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(9)
    doc.core_properties.title = DATA['title']
    doc.core_properties.author = 'Team Doom'
    doc.core_properties.subject = 'Stage 2 beginner test bank'
    doc.core_properties.comments = 'Draft for teammate review'
    doc.add_paragraph(DATA['title'], 'Title')
    doc.add_paragraph('Team Doom  |  Version 1.0 draft  |  30 October 2026', 'Subtitle')
    doc.add_paragraph('Practise the main ideas in Hardware, Networks and AI Behaviour. All 46 questions can be answered without equipment or an external website. Write answers separately, then use the answer key to check your understanding.')
    doc.add_paragraph('Draft for teammate review. Topic labels refer to the shared learning material. Final Study Guide section references will be added when the guide is available.')
    p = doc.add_paragraph('Sources: Hardware and Beginner Lab Skills; Network Fundamentals and Beginner Prerequisites; AI Behaviour and References. These chapters are in Learning Materials, grouped by subject.')
    for run in p.runs:
        run.font.size = Pt(9)

    def paragraph(text, bold=False, small=False, keep=False, indent=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if small:
            run.font.size = Pt(9)
        p.paragraph_format.keep_with_next = keep
        p.paragraph_format.keep_together = True
        if indent:
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(2)
        return p

    for prefix, title, instruction in [
        ('TF', 'Part A True or False', 'Write True or False for each statement.'),
        ('MC', 'Part B Multiple Choice', 'Choose ONE answer, A, B, C or D, for each question.'),
        ('OE', 'Part C Explain Your Answer', 'Answer all three parts of each question. Aim for a developed response, normally more than 100 words in total per question, or equivalent reasoning through a labeled diagram, table or step-by-step trace. Keep your explanation simple and use the supplied situation.')]:
        heading = doc.add_heading(title, 1)
        heading.paragraph_format.page_break_before = prefix != 'TF'
        paragraph(instruction)
        topic = None
        for q in (q for q in QUESTIONS if q['id'].startswith(prefix)):
            if q['id'] in ('MC08', 'MC15', 'OE04'):
                topic = None
            if topic != q['topic']:
                topic = q['topic']
                heading = doc.add_heading(topic, 2)
                heading.paragraph_format.page_break_before = q['id'] in ('MC08', 'MC15', 'OE04')
            stem = paragraph(f"{q['id']}  {q['prompt']}", keep=True)
            if prefix == 'MC':
                stem.paragraph_format.space_after = Pt(3)
                stem.paragraph_format.line_spacing = 1.04
                for letter, option in zip(LETTERS, q['options']):
                    choice = paragraph(f'{letter}. {option}', keep=True, indent=True)
                    choice.paragraph_format.line_spacing = 1.0
                    choice.paragraph_format.space_after = Pt(0)
            elif prefix == 'OE':
                for letter, task in zip('abc', q['tasks']):
                    paragraph(f'{letter}) {task}', keep=True)
            p = paragraph('Topic source  ' + q['source'], small=True)
            p.paragraph_format.space_after = Pt(9)
    heading = doc.add_heading('Answer Key', 1)
    heading.paragraph_format.page_break_before = True
    paragraph('Try the questions before reading this section. Equivalent correct explanations, labeled diagrams and worked steps are acceptable for open-ended answers.')
    for prefix, title in [('TF', 'True or False Answers'), ('MC', 'Multiple Choice Answers'), ('OE', 'Open Ended Answers')]:
        heading = doc.add_heading(title, 2)
        heading.paragraph_format.page_break_before = prefix != 'TF'
        for q in (q for q in QUESTIONS if q['id'].startswith(prefix)):
            if prefix == 'OE':
                heading = doc.add_heading(q['id'] + ' Sample Answer', 3)
                heading.paragraph_format.page_break_before = q['id'] == 'OE04'
                paragraph(q['sample'])
                paragraph('Expected key points', bold=True, keep=True)
                for i, point in enumerate(q['key_points']):
                    paragraph('• ' + point, keep=i < len(q['key_points']) - 1)
            else:
                answer = q['answer']
                if prefix == 'MC':
                    answer += '  ' + q['options'][LETTERS.index(q['answer'])]
                p = doc.add_paragraph()
                p.paragraph_format.page_break_before = q['id'] == 'MC11'
                p.paragraph_format.keep_together = True
                p.add_run(f"{q['id']}  {answer}. ").bold = True
                p.add_run(q['explanation'])
    target = ROOT / 'Doom_Stage2_TestBank_V1.docx'
    doc.save(target)
    # Reopen the package and verify each question and answer label survived export.
    reopened = Document(target)
    text = '\n'.join(p.text for p in reopened.paragraphs)
    assert text.index('Answer Key') > text.index('OE06  ')
    for q in QUESTIONS:
        assert q['prompt'] in text
        assert (q.get('sample') or q['explanation']) in text
    print('Created Questions.md, AnswerKey.md and editable Word bank with all questions before the answer key.')


if __name__ == '__main__':
    check()
    if '--check' not in sys.argv:
        markdown()
        if '--markdown-only' not in sys.argv:
            word()
